#!/usr/bin/env python3
"""Generate the trading-card inventory portfolio management workbook.

Builds 在庫ポートフォリオ管理.xlsx with 6 sheets:
使い方 / 設定 / ウォッチリスト / 仕入れ計画 / 在庫管理 / ダッシュボード

Scope: トレーディングカードのみ（ポケモンカード／ワンピースカード／遊戯王 等）。
4象限（流動性×収益性）は想定売却日数・想定粗利率から自動判定する。
各象限で「仕入れてよいか」の判断基準は 希少性・トレンド・ボラティリティ の3条件。

Run: python3 build_workbook.py
Then: python3 <xlsx-skill>/scripts/recalc.py ../在庫ポートフォリオ管理.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.chart import PieChart, BarChart, Reference
import os

OUT_PATH = os.path.join(os.path.dirname(__file__), "..", "在庫ポートフォリオ管理.xlsx")

FONT_NAME = "Yu Gothic"

BLUE = Font(name=FONT_NAME, size=10, color="0000FF")
BLACK = Font(name=FONT_NAME, size=10, color="000000")
GREEN = Font(name=FONT_NAME, size=10, color="008000")
BOLD = Font(name=FONT_NAME, size=10, bold=True)
BOLD_WHITE = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
TITLE = Font(name=FONT_NAME, size=16, bold=True, color="1F4E78")
SECTION = Font(name=FONT_NAME, size=11, bold=True, color="1F4E78")
NOTE = Font(name=FONT_NAME, size=9, italic=True, color="666666")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CUR = '¥#,##0;[RED]-¥#,##0;"-"'
PCT = '0.0%;[RED]-0.0%;"-"'
INT = '#,##0;[RED]-#,##0;"-"'
DATE = "yyyy/mm/dd"

wb = openpyxl.Workbook()
wb.remove(wb.active)


def sheet(name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    return ws


def title(ws, cell, text):
    ws[cell] = text
    ws[cell].font = TITLE


def note(ws, cell, text):
    ws[cell] = text
    ws[cell].font = NOTE


def section(ws, cell, text):
    ws[cell] = text
    ws[cell].font = SECTION
    ws[cell].fill = SECTION_FILL


def header_row(ws, row, col_start, labels):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col_start + i, value=label)
        c.font = BOLD_WHITE
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def cols(names):
    """Map column name -> letter, in order, starting at column A."""
    return {name: get_column_letter(i + 1) for i, name in enumerate(names)}


# =====================================================================
# 1. 使い方
# =====================================================================
ws = sheet("使い方")
set_col_widths(ws, [4, 100])
title(ws, "B2", "トレカ在庫ポートフォリオ管理ブック")
note(ws, "B3", "スニダン等で扱うトレーディングカード（ポケモンカード／ワンピースカード／遊戯王 等）の在庫を4象限で管理するための運用テンプレートです。")

rows = [
    ("■ シート構成", None),
    ("設定", "4象限のしきい値、銘柄タイトル／希少性／トレンド／ボラティリティ／ステータスのマスタ、集中リスク集計用のキャラ・イラストレーターリストを管理します。まずこのシートを事業実態に合わせて調整してください。"),
    ("ウォッチリスト", "スニダン等で気になっている未仕入れのカードを登録すると、想定売却日数と想定粗利率から自動で4象限に振り分けます。"),
    ("仕入れ計画", "象限ごとに「何を・いくらまで・どの希少性／トレンド／ボラティリティなら仕入れてよいか」のルールと月間予算配分を定義します。"),
    ("在庫管理", "実在庫を4分類＋タグ（キャラ／イラストレーター／絶版／周年／希少性／トレンド／ボラティリティ）で登録し、ポートフォリオ比率を自動計算します。"),
    ("ダッシュボード", "資金効率（在庫の入れ替わり回数・資金倍率・含み損益率）と換金のしやすさ、偏りリスクを自動集計します。"),
    ("■ 2段階の判断の仕組み", None),
    ("① 自動判定（象限）", "「想定売却日数」と「想定粗利率」の2つの数値から、コア／回転／厳選／見送りの4象限を自動で判定します（「設定」シートのしきい値で調整可能）。"),
    ("② 仕入れ可否の判断（属性）", "象限が決まったら、そのカードの「希少性」「トレンド（値動きの勢い）」「ボラティリティ（値動きの荒さ）」が、その象限で仕入れてよい条件に当てはまるかを「仕入れ計画」シートと照らし合わせて判断します。"),
    ("■ 入力ルール（色分け）", None),
    ("青字", "手入力セル。実データに置き換えてください。"),
    ("黄色背景", "しきい値・予算など、事業判断が必要な重要な前提値。"),
    ("黒字", "自動計算セル。数式が入っているため直接編集しないでください。"),
    ("緑字", "他シートを参照しているセル。"),
    ("■ 使い始め方", None),
    ("Step1", "「設定」シートで銘柄タイトル（ポケカ／ワンピカ／遊戯王 等）・希少性・4象限しきい値を確認・調整する。"),
    ("Step2", "「ウォッチリスト」にスニダンのウォッチリスト登録分を入力し、象限分布を確認する。"),
    ("Step3", "「仕入れ計画」で象限別の上限単価・許容希少性／トレンド／ボラティリティ・予算配分を確定する。"),
    ("Step4", "「在庫管理」に実在庫を棚卸し登録し、ポートフォリオ比率と偏りリスクを確認する。"),
    ("Step5", "「ダッシュボード」を週次で確認し、仕入れ計画にフィードバックする。"),
    ("■ 用語解説（投資フレームワーク検討資料に準拠）", None),
    ("回転率", "一定期間内に在庫が何回入れ替わる（売れる）かを示す指標。本ブックでは 365日÷想定売却日数 で近似します。"),
    ("資本効率", "投じた資金に対してどれだけ効率よく利益を生めるか。「想定粗利率×回転率」で近似し、ウォッチリスト・在庫管理シートに列を設けています。同じ象限内でカードを比較する際、この値が高いものを優先的に仕入れる基準になります。"),
    ("実効流動性", "理論上の流動性（想定売却日数）から、ボラティリティによる値崩れリスクを差し引いた、実質的な換金力。本ブックでは「実効想定売却日数＝想定売却日数×ボラティリティ係数」で近似します（設定シートで係数を調整可能）。"),
    ("ポジションサイズ目安", "「期待リターン（想定粗利率）÷ボラティリティ」で近似した、その銘柄をどれだけ厚く持ってよいかの目安スコア。同じ象限内でも、ボラティリティが高いものほど保有量を抑える判断に使います。"),
    ("相関ドライバー", "複数の商材が同じ要因（キャラクター人気・絵師人気・収録弾の再販発表・周年企画等）で連動して値動きする際の共通要因。ジャンル（銘柄タイトル）が分散していても、同じキャラ・イラストレーター・企画に在庫が偏っていると、その人気が落ちた瞬間に在庫全体が一緒に値下がりするリスクがあります。ダッシュボードでキャラ別・イラストレーター別・企画/収録弾別・希少性帯別に偏りをチェックできます。"),
    ("投資フレームワークとの対応", "本ブックの「コア／回転／厳選／見送り」は、社内の投資フレームワーク検討資料でいう「コアポジション／キャッシュ同等物／グロース枠／劣後在庫」にそれぞれ対応します（詳細は「設定」シートの4象限マスタ参照）。"),
    ("■ 注意事項", None),
    ("サンプル行", "各シートの数行はサンプル値です。実データ入力時は上書き、または不要なら内容を削除してください（数式は残しても構いません）。"),
    ("偏りリスク(HHI)", "特定のカテゴリ・キャラ・イラストレーターに在庫が偏っていないかを数値化したものです（集中度指数／HHI）。数値が大きいほど偏りが大きいことを意味します。目安：1,500未満＝分散、1,500〜2,500＝中程度、2,500以上＝偏りが大きい。"),
]

r = 5
for label, desc in rows:
    if desc is None:
        section(ws, f"B{r}", label)
        r += 1
        continue
    ws[f"B{r}"] = f"　{label}：{desc}"
    ws[f"B{r}"].font = BLACK
    ws[f"B{r}"].alignment = Alignment(wrap_text=True)
    r += 1
ws.row_dimensions[2].height = 24

# =====================================================================
# 2. 設定
# =====================================================================
ws = sheet("設定")
set_col_widths(ws, [3, 46, 16, 16, 40, 20, 14])
title(ws, "B2", "設定")
note(ws, "B3", "4象限のしきい値とマスタデータ。事業実態に合わせて調整してください。")

section(ws, "B5", "■ 4象限 自動判定しきい値（想定売却日数・想定粗利率から象限を機械的に判定）")
ws["B6"] = "早く売れる基準（想定売却日数・この日数以内なら「早く売れる」とみなす）"
ws["C6"] = 30
ws["C6"].font, ws["C6"].fill, ws["C6"].number_format = BLUE, INPUT_FILL, '0"日"'
ws["B7"] = "儲かる基準（想定粗利率・この率以上なら「利益率が高い」とみなす）"
ws["C7"] = 0.20
ws["C7"].font, ws["C7"].fill, ws["C7"].number_format = BLUE, INPUT_FILL, PCT
ws["B8"] = "売れ残りリスク基準（想定売却日数がこの日数を超えたら要注意）"
ws["C8"] = 90
ws["C8"].font, ws["C8"].fill, ws["C8"].number_format = BLUE, INPUT_FILL, '0"日"'
for rr in (6, 7, 8):
    ws[f"B{rr}"].font = BLACK
    ws[f"B{rr}"].border = ws[f"C{rr}"].border = BORDER
note(ws, "B9", "※ この2つのしきい値は象限の自動判定専用です。カードごとの仕入れ可否は「仕入れ計画」シートの希少性／トレンド／ボラティリティ条件で判断します。")

section(ws, "B11", "■ 4象限マスタ（分類ロジック・方針の目安・投資フレームワークとの対応）")
header_row(ws, 12, 2, ["象限コード", "早く売れる？", "利益率は高い？", "方針の目安", "投資フレームワークでの呼称", "資金配分の目安"])
quad_master = [
    ("I(コア)", "早い", "高い", "積極的に仕入れる中心銘柄。予算配分と在庫上限を厚めに。", "コアポジション（収益基盤）", 0.55),
    ("II(回転)", "早い", "低い（薄利）", "資金回転を稼ぐ数量枠。単価上限を抑えて回転数で稼ぐ。", "キャッシュ同等物（リスクバッファ）", 0.20),
    ("III(厳選)", "遅い", "高い", "資金拘束を許容できる範囲で少数厳選。上限点数を絞る。", "グロース枠（将来リターン）", 0.20),
    ("IV(見送り)", "遅い", "低い（薄利）", "原則仕入れ見送り。ウォッチのみ継続し再評価を待つ。", "劣後在庫（排除対象）", 0.05),
]
r = 13
for row in quad_master:
    for i, v in enumerate(row):
        c = ws.cell(row=r, column=2 + i, value=v)
        c.font = BLACK
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if i == 5:
            c.number_format = PCT
    r += 1
note(ws, "B17", "※ 「投資フレームワークでの呼称」「資金配分の目安」は社内の投資フレームワーク検討資料（コア・サテライト型：コアが過半・劣後は最小限）を参考にした値です。「仕入れ計画」シートの予算配分比率の初期値として反映しています。")

section(ws, "B18", "■ 銘柄タイトルマスタ（トレカのタイトル。在庫管理・ウォッチリストのプルダウンで使用）")
header_row(ws, 19, 2, ["銘柄タイトル"])
titles = ["ポケモンカード", "ワンピースカード", "遊戯王", "デュエル・マスターズ", "マジック：ザ・ギャザリング", "その他"]
r = 20
for t in titles:
    ws.cell(row=r, column=2, value=t).font = BLUE
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r, column=2).border = BORDER
    r += 1
CAT_RANGE = "設定!$B$20:$B$25"

section(ws, "B27", "■ 希少性マスタ（レアリティ。SSが最も希少）")
header_row(ws, 28, 2, ["希少性"])
rarities = ["SS", "S", "A", "B", "C"]
r = 29
for rar in rarities:
    ws.cell(row=r, column=2, value=rar).font = BLUE
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r, column=2).border = BORDER
    r += 1
RARITY_RANGE = "設定!$B$29:$B$33"

section(ws, "B35", "■ トレンドマスタ（価格・需要の勢い）")
header_row(ws, 36, 2, ["トレンド"])
trends = ["上昇", "安定", "下降"]
r = 37
for t in trends:
    ws.cell(row=r, column=2, value=t).font = BLUE
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r, column=2).border = BORDER
    r += 1
TREND_RANGE = "設定!$B$37:$B$39"

section(ws, "B41", "■ ボラティリティマスタ（価格の変動の荒さ・実効流動性の計算に使用）")
header_row(ws, 42, 2, ["ボラティリティ", "目減り係数"])
vols = [("高", 1.5), ("中", 1.2), ("低", 1.0)]
r = 43
for v, factor in vols:
    ws.cell(row=r, column=2, value=v).font = BLUE
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r, column=2).border = BORDER
    ws.cell(row=r, column=3, value=factor).font = BLUE
    ws.cell(row=r, column=3).fill = INPUT_FILL
    ws.cell(row=r, column=3).border = BORDER
    ws.cell(row=r, column=3).number_format = "0.00"
    r += 1
VOL_RANGE = "設定!$B$43:$B$45"
VOL_FACTOR_RANGE = "設定!$C$43:$C$45"
note(ws, "B46", "※ 目減り係数：想定売却日数にこの係数を掛けて「実効想定売却日数」を算出します（値動きが荒いほど、安全に売り切るまでに実質的な時間がかかるとみなす）。")

section(ws, "B47", "■ ステータスマスタ")
header_row(ws, 48, 2, ["ステータス"])
statuses = ["在庫中", "受注済", "売却済"]
r = 49
for st in statuses:
    ws.cell(row=r, column=2, value=st).font = BLUE
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r, column=2).border = BORDER
    r += 1
STATUS_RANGE = "設定!$B$49:$B$51"

section(ws, "B53", "■ キャラクター集中リスク集計対象リスト（在庫管理シートで使うキャラ名を入力すると、ダッシュボードで自動集計されます。最大20件）")
header_row(ws, 54, 2, ["キャラ名"])
sample_chars = ["リザードン", "ルフィ", "青眼の白龍"]
CHAR_START = 55
for i in range(20):
    r = CHAR_START + i
    v = sample_chars[i] if i < len(sample_chars) else ""
    ws.cell(row=r, column=2, value=v).font = BLUE
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r, column=2).border = BORDER
CHAR_RANGE = f"設定!$B${CHAR_START}:$B${CHAR_START+19}"

section(ws, "B76", "■ イラストレーター集中リスク集計対象リスト（最大20件）")
header_row(ws, 77, 2, ["イラストレーター名"])
sample_illust = ["サンプル絵師A", "サンプル絵師B"]
ILLUST_START = 78
for i in range(20):
    r = ILLUST_START + i
    v = sample_illust[i] if i < len(sample_illust) else ""
    ws.cell(row=r, column=2, value=v).font = BLUE
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r, column=2).border = BORDER
ILLUST_RANGE = f"設定!$B${ILLUST_START}:$B${ILLUST_START+19}"

section(ws, "B100", "■ 企画・収録弾 集中リスク集計対象リスト（相関ドライバー。収録弾／レアリティ帯／周年企画など、キャラ・イラストレーター以外で値動きが連動する要因。在庫管理シートで使う値を入力すると、ダッシュボードで自動集計されます。最大20件）")
header_row(ws, 101, 2, ["企画・収録弾名"])
sample_campaigns = ["スカーレット&バイオレット 151", "ROMANCE DAWN", "25周年記念"]
CAMPAIGN_START = 102
for i in range(20):
    r = CAMPAIGN_START + i
    v = sample_campaigns[i] if i < len(sample_campaigns) else ""
    ws.cell(row=r, column=2, value=v).font = BLUE
    ws.cell(row=r, column=2).fill = INPUT_FILL
    ws.cell(row=r, column=2).border = BORDER
CAMPAIGN_RANGE = f"設定!$B${CAMPAIGN_START}:$B${CAMPAIGN_START+19}"
note(ws, f"B{CAMPAIGN_START+21}",
     "※ 相関ドライバーとは、複数の商材が同じ要因（人気投票・映画出演・再販発表・環境トップ入り等）で連動して値動きする際の共通要因です。"
     "ジャンル（銘柄タイトル）が分散していても、同じ収録弾・企画に在庫が偏っていると、その人気が落ちた瞬間に在庫全体が一緒に値下がりするリスクがあります。")

print("設定 sheet done:", CAT_RANGE, RARITY_RANGE, TREND_RANGE, VOL_RANGE, STATUS_RANGE, CHAR_RANGE, ILLUST_RANGE, CAMPAIGN_RANGE)

# =====================================================================
# 3. ウォッチリスト
# =====================================================================
ws = sheet("ウォッチリスト")
WL_COLS = ["No", "銘柄名", "銘柄タイトル", "希少性", "トレンド", "ボラティリティ",
           "スニダン参考価格", "想定仕入値", "想定売却価格", "想定売却日数", "実効想定売却日数",
           "想定粗利額", "想定粗利率", "資本効率", "ポジションサイズ目安", "象限", "仕入れ判定", "メモ"]
WC = cols(WL_COLS)
set_col_widths(ws, [5, 24, 16, 10, 10, 12, 13, 12, 12, 12, 14, 12, 12, 12, 14, 12, 16, 24])
title(ws, "B2", "ウォッチリスト（4象限分類）")
note(ws, "B3", "スニダン等のウォッチリスト登録分を入力すると、想定売却日数・想定粗利率から自動で4象限に振り分けます。青字セルのみ入力してください。")

section(ws, "B5", "■ 象限別 候補数サマリー")
labels = ["I(コア) 候補数", "II(回転) 候補数", "III(厳選) 候補数", "IV(見送り) 候補数"]
WL_HEADER_ROW = 12
WL_DATA_START = 13
WL_DATA_END = WL_DATA_START + 99  # 100 rows
for i, lab in enumerate(labels):
    r = 6 + i
    ws[f"B{r}"] = lab
    ws[f"B{r}"].font = BLACK
    ws[f"B{r}"].border = BORDER
    quad_code = ["I(コア)", "II(回転)", "III(厳選)", "IV(見送り)"][i]
    ws[f"C{r}"] = f'=COUNTIF({WC["象限"]}{WL_DATA_START}:{WC["象限"]}{WL_DATA_END},"{quad_code}")'
    ws[f"C{r}"].font = BLACK
    ws[f"C{r}"].number_format = INT
    ws[f"C{r}"].border = BORDER

wl_header_labels = ["No", "銘柄名", "銘柄タイトル", "希少性", "トレンド", "ボラティリティ",
                     "スニダン参考価格\n（現在相場）", "想定仕入値", "想定売却価格",
                     "想定売却\n日数", "実効想定売却\n日数（ボラ考慮）", "想定粗利額", "想定粗利率",
                     "資本効率\n（利益率×回転率）", "ポジション\nサイズ目安", "象限", "仕入れ判定", "メモ"]
header_row(ws, WL_HEADER_ROW, 1, wl_header_labels)
ws.freeze_panes = f"A{WL_DATA_START}"

wl_samples = [
    ("リザードンVMAX（ポケモンカード）", "ポケモンカード", "SS", "上昇", "低", 28000, 22000, 32000, 15, ""),
    ("ルフィ SEC（ワンピースカード）", "ワンピースカード", "S", "上昇", "中", 18000, 15000, 19500, 20, ""),
    ("青眼の白龍 初期（遊戯王）", "遊戯王", "SS", "安定", "低", 45000, 38000, 52000, 60, "絶版・入手困難"),
    ("ピカチュウ プロモ（ポケモンカード）", "ポケモンカード", "A", "安定", "中", 9000, 7000, 8500, 45, "周年記念"),
    ("人気シングル（遊戯王）", "遊戯王", "A", "上昇", "中", 12000, 10000, 11500, 10, ""),
    ("一般パック品（遊戯王）", "遊戯王", "C", "下降", "高", 3000, 2800, 2900, 90, ""),
]
wl_sample_names = ["銘柄名", "銘柄タイトル", "希少性", "トレンド", "ボラティリティ",
                    "スニダン参考価格", "想定仕入値", "想定売却価格", "想定売却日数", "メモ"]
for i, values in enumerate(wl_samples):
    r = WL_DATA_START + i
    for name, val in zip(wl_sample_names, values):
        ws[f"{WC[name]}{r}"] = val

for r in range(WL_DATA_START, WL_DATA_END + 1):
    B, H, I_, J, K, L, M, N = (WC["銘柄名"], WC["想定仕入値"], WC["想定売却価格"], WC["想定売却日数"],
                               WC["想定粗利額"], WC["想定粗利率"], WC["象限"], WC["仕入れ判定"])
    Vc, Eff, Cap, Pos = WC["ボラティリティ"], WC["実効想定売却日数"], WC["資本効率"], WC["ポジションサイズ目安"]
    ws[f"{WC['No']}{r}"] = f'=IF({B}{r}="","",ROW()-{WL_DATA_START-1})'
    ws[f"{K}{r}"] = f'=IF({B}{r}="","",{I_}{r}-{H}{r})'
    ws[f"{L}{r}"] = f'=IF({B}{r}="","",IF({H}{r}=0,"",{K}{r}/{H}{r}))'
    ws[f"{Eff}{r}"] = (f'=IF({B}{r}="","",{J}{r}*IFERROR(INDEX({VOL_FACTOR_RANGE},MATCH({Vc}{r},{VOL_RANGE},0)),1))')
    ws[f"{Cap}{r}"] = f'=IF({B}{r}="","",IF({J}{r}=0,"",{L}{r}*365/{J}{r}))'
    ws[f"{Pos}{r}"] = (f'=IF({B}{r}="","",IFERROR({L}{r}/INDEX({VOL_FACTOR_RANGE},MATCH({Vc}{r},{VOL_RANGE},0)),""))')
    ws[f"{M}{r}"] = (f'=IF({B}{r}="","",IF({J}{r}<=設定!$C$6,IF({L}{r}>=設定!$C$7,"I(コア)","II(回転)"),'
                      f'IF({L}{r}>=設定!$C$7,"III(厳選)","IV(見送り)")))')
    ws[f"{N}{r}"] = (f'=IF({B}{r}="","",IF(OR({M}{r}="I(コア)",{M}{r}="II(回転)"),"仕入れ候補",'
                      f'IF({M}{r}="III(厳選)","厳選検討（少数）","見送り")))')
    col_font_fmt = {
        "No": (BLACK, INT), "銘柄名": (BLUE, None), "銘柄タイトル": (BLUE, None), "希少性": (BLUE, None),
        "トレンド": (BLUE, None), "ボラティリティ": (BLUE, None), "スニダン参考価格": (BLUE, CUR),
        "想定仕入値": (BLUE, CUR), "想定売却価格": (BLUE, CUR), "想定売却日数": (BLUE, '0"日"'),
        "実効想定売却日数": (BLACK, '0.0"日"'), "想定粗利額": (BLACK, CUR), "想定粗利率": (BLACK, PCT),
        "資本効率": (BLACK, PCT), "ポジションサイズ目安": (BLACK, "0.00"), "象限": (BLACK, None),
        "仕入れ判定": (BLACK, None), "メモ": (BLUE, None),
    }
    for name, (font, fmt) in col_font_fmt.items():
        c = ws[f"{WC[name]}{r}"]
        c.font = font
        c.border = BORDER
        if fmt:
            c.number_format = fmt

dv_cat = DataValidation(type="list", formula1=CAT_RANGE, allow_blank=True)
dv_rarity_wl = DataValidation(type="list", formula1=RARITY_RANGE, allow_blank=True)
dv_trend_wl = DataValidation(type="list", formula1=TREND_RANGE, allow_blank=True)
dv_vol_wl = DataValidation(type="list", formula1=VOL_RANGE, allow_blank=True)
for dv in (dv_cat, dv_rarity_wl, dv_trend_wl, dv_vol_wl):
    ws.add_data_validation(dv)
dv_cat.add(f"{WC['銘柄タイトル']}{WL_DATA_START}:{WC['銘柄タイトル']}{WL_DATA_END}")
dv_rarity_wl.add(f"{WC['希少性']}{WL_DATA_START}:{WC['希少性']}{WL_DATA_END}")
dv_trend_wl.add(f"{WC['トレンド']}{WL_DATA_START}:{WC['トレンド']}{WL_DATA_END}")
dv_vol_wl.add(f"{WC['ボラティリティ']}{WL_DATA_START}:{WC['ボラティリティ']}{WL_DATA_END}")

print("ウォッチリスト sheet done")

# 在庫管理シートの列レイアウトは仕入れ計画から先に参照するため、ここで定義しておく
# （実際のシート内容の作成は「5. 在庫管理」セクションで行う）
INV_COLS = ["No", "商品名", "銘柄タイトル", "キャラ", "イラストレーター", "企画・収録弾", "絶版", "周年",
            "希少性", "トレンド", "ボラティリティ", "ステータス", "仕入日", "仕入値",
            "現在時価", "想定売却日数", "実効想定売却日数", "想定粗利額", "想定粗利率",
            "資本効率", "ポジションサイズ目安", "象限", "ポートフォリオ比率", "メモ"]
IC = cols(INV_COLS)

# =====================================================================
# 4. 仕入れ計画
# =====================================================================
ws = sheet("仕入れ計画")
PLAN_COLS = ["象限", "呼称", "方針", "対象タイトル例", "上限仕入単価",
             "許容希少性", "許容トレンド", "許容ボラティリティ",
             "予算配分比率", "予算配分金額", "在庫上限点数", "ウォッチ候補数", "現在庫点数"]
PC = cols(PLAN_COLS)
set_col_widths(ws, [10, 10, 26, 20, 13, 12, 12, 14, 11, 13, 11, 11, 11])
title(ws, "B2", "仕入れ計画（象限別ルール）")
note(ws, "B3", "「何を・いくらまで・どの希少性／トレンド／ボラティリティなら仕入れてよいか」を象限ごとに定義し、月間予算を配分します。黄色セルが判断の起点です。")

INV_HEADER_ROW = 10
INV_DATA_START = 11
INV_DATA_END = INV_DATA_START + 99  # 100 rows

ws["B5"] = "月間仕入れ総予算"
ws["B5"].font = BLACK
ws["C5"] = 500000
ws["C5"].font, ws["C5"].fill, ws["C5"].number_format = BLUE, INPUT_FILL, CUR
ws["B5"].border = ws["C5"].border = BORDER

PLAN_HEADER_ROW = 7
PLAN_START = 8
header_row(ws, PLAN_HEADER_ROW, 1, PLAN_COLS)

plan_rows = [
    ("I(コア)", "I(コア)", "積極的に仕入れる中心銘柄。優先確保。", "ポケモンカード、ワンピースカード", 50000,
     "S以上", "上昇", "低〜中", 0.55, 15),
    ("II(回転)", "II(回転)", "資金回転重視。数量で稼ぐ。", "遊戯王、デュエル・マスターズ", 15000,
     "A以上", "上昇・安定", "指定なし", 0.20, 40),
    ("III(厳選)", "III(厳選)", "少数厳選。資金拘束を許容できる範囲で。", "ポケモンカード、ワンピースカード", 80000,
     "SS", "上昇・安定", "低", 0.20, 5),
    ("IV(見送り)", "IV(見送り)", "原則仕入れ見送り。ウォッチのみ継続。", "-", 0,
     "-", "下降", "高", 0.05, 0),
]
for i, (code, label, policy, titles_ex, max_price, rarity, trend, vol, ratio, cap) in enumerate(plan_rows):
    r = PLAN_START + i
    ws[f"{PC['象限']}{r}"] = code
    ws[f"{PC['呼称']}{r}"] = label
    ws[f"{PC['方針']}{r}"] = policy
    ws[f"{PC['対象タイトル例']}{r}"] = titles_ex
    ws[f"{PC['上限仕入単価']}{r}"] = max_price
    ws[f"{PC['許容希少性']}{r}"] = rarity
    ws[f"{PC['許容トレンド']}{r}"] = trend
    ws[f"{PC['許容ボラティリティ']}{r}"] = vol
    ws[f"{PC['予算配分比率']}{r}"] = ratio
    ws[f"{PC['予算配分金額']}{r}"] = f"=$C$5*{PC['予算配分比率']}{r}"
    ws[f"{PC['在庫上限点数']}{r}"] = cap
    ws[f"{PC['ウォッチ候補数']}{r}"] = f'=COUNTIF(ウォッチリスト!${WC["象限"]}${WL_DATA_START}:${WC["象限"]}${WL_DATA_END},"{code}")'
    ws[f"{PC['現在庫点数']}{r}"] = f'=COUNTIF(在庫管理!${IC["象限"]}${INV_DATA_START}:${IC["象限"]}${INV_DATA_END},"{code}")'
    fmt_map = {
        "象限": (BLACK, None, False), "呼称": (BLACK, None, False), "方針": (BLUE, None, False),
        "対象タイトル例": (BLUE, None, False), "上限仕入単価": (BLUE, CUR, True),
        "許容希少性": (BLUE, None, True), "許容トレンド": (BLUE, None, True), "許容ボラティリティ": (BLUE, None, True),
        "予算配分比率": (BLUE, PCT, True), "予算配分金額": (BLACK, CUR, False),
        "在庫上限点数": (BLUE, INT, True), "ウォッチ候補数": (BLACK, INT, False), "現在庫点数": (BLACK, INT, False),
    }
    for name, (font, fmt, fill) in fmt_map.items():
        c = ws[f"{PC[name]}{r}"]
        c.font = font
        c.border = BORDER
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = INPUT_FILL
        if name in ("方針", "対象タイトル例"):
            c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 42

TOTAL_ROW = PLAN_START + len(plan_rows)
ws[f"{PC['呼称']}{TOTAL_ROW}"] = "合計"
ws[f"{PC['呼称']}{TOTAL_ROW}"].font = BOLD
ws[f"{PC['予算配分比率']}{TOTAL_ROW}"] = f"=SUM({PC['予算配分比率']}{PLAN_START}:{PC['予算配分比率']}{TOTAL_ROW-1})"
ws[f"{PC['予算配分金額']}{TOTAL_ROW}"] = f"=SUM({PC['予算配分金額']}{PLAN_START}:{PC['予算配分金額']}{TOTAL_ROW-1})"
for name in ("呼称", "予算配分比率", "予算配分金額"):
    c = ws[f"{PC[name]}{TOTAL_ROW}"]
    c.font = BOLD
    c.border = BORDER
    if name == "予算配分比率":
        c.number_format = PCT
    if name == "予算配分金額":
        c.number_format = CUR

ws.conditional_formatting.add(
    f"{PC['予算配分比率']}{TOTAL_ROW}",
    CellIsRule(operator="notEqual", formula=["1"], fill=PatternFill("solid", fgColor="FFC7CE"),
               font=Font(color="9C0006")),
)
note(ws, f"B{TOTAL_ROW+2}", "※ 予算配分比率の合計は100%になるように調整してください（赤色表示は100%以外を意味します）。")
note(ws, f"B{TOTAL_ROW+3}", "※ 許容希少性／許容トレンド／許容ボラティリティは、ウォッチリスト・在庫管理シートの各カードの値と見比べて仕入れ可否を判断するための基準です。")
note(ws, f"B{TOTAL_ROW+4}",
     "※ 予算配分の初期値（55%/20%/20%/5%）はコア・サテライト型投資の考え方（コアが過半・劣後は最小限）を参考にした値です。"
     "同じ象限内でカードを選ぶ際は、ウォッチリスト・在庫管理シートの「資本効率」が高いものを優先し、「ポジションサイズ目安」が低いもの（＝ボラティリティが高いもの）は保有量を絞ってください。")

print("仕入れ計画 sheet done")

# =====================================================================
# 5. 在庫管理
# =====================================================================
ws = sheet("在庫管理")
# INV_COLS / IC は「4. 仕入れ計画」セクションの手前で定義済み
set_col_widths(ws, [5, 26, 14, 12, 14, 16, 8, 8, 9, 9, 12, 10, 12, 11, 12, 11, 14, 11, 11, 11, 14, 11, 12, 22])
title(ws, "B2", "在庫管理シート（実在庫・ポートフォリオ管理）")
note(ws, "B3", "青字セルのみ入力してください。ステータスが「在庫中」の行のみ、集計・ポートフォリオ比率に反映されます。")

summary_labels = ["総仕入額（在庫中）", "総時価（在庫中）", "含み損益", "含み損益率", "在庫点数（在庫中）"]
for i, lab in enumerate(summary_labels):
    r = 5 + i
    ws.cell(row=r, column=2, value=lab).font = BLACK
    ws.cell(row=r, column=2).border = BORDER
K, N, ST = IC["仕入値"], IC["現在時価"], IC["ステータス"]
ws["C5"] = f'=SUMIFS({K}{INV_DATA_START}:{K}{INV_DATA_END},{ST}{INV_DATA_START}:{ST}{INV_DATA_END},"在庫中")'
ws["C6"] = f'=SUMIFS({N}{INV_DATA_START}:{N}{INV_DATA_END},{ST}{INV_DATA_START}:{ST}{INV_DATA_END},"在庫中")'
ws["C7"] = "=C6-C5"
ws["C8"] = '=IFERROR(C7/C5,"")'
ws["C9"] = f'=COUNTIF({ST}{INV_DATA_START}:{ST}{INV_DATA_END},"在庫中")'
for r, fmt in [(5, CUR), (6, CUR), (7, CUR), (8, PCT), (9, INT)]:
    ws.cell(row=r, column=3).font = BLACK
    ws.cell(row=r, column=3).border = BORDER
    ws.cell(row=r, column=3).number_format = fmt

inv_header_labels = ["No", "商品名", "銘柄タイトル", "キャラ", "イラストレーター", "企画・収録弾\n（相関ドライバー）",
                      "絶版", "周年", "希少性", "トレンド", "ボラティリティ", "ステータス", "仕入日", "仕入値",
                      "現在時価\n（想定売却価格）", "想定売却\n日数", "実効想定売却\n日数（ボラ考慮）",
                      "想定粗利額", "想定粗利率", "資本効率\n（利益率×回転率）", "ポジション\nサイズ目安",
                      "象限", "ポートフォリオ\n比率", "メモ"]
header_row(ws, INV_HEADER_ROW, 1, inv_header_labels)
ws.freeze_panes = f"A{INV_DATA_START}"

inv_samples = [
    ("リザードンVMAX（ポケモンカード）", "ポケモンカード", "リザードン", "サンプル絵師A", "スカーレット&バイオレット 151",
     "×", "×", "SS", "上昇", "低", "在庫中", "2026-06-15", 24000, 31000, 18, ""),
    ("ルフィ SEC（ワンピースカード）", "ワンピースカード", "ルフィ", "", "ROMANCE DAWN",
     "×", "×", "S", "上昇", "中", "在庫中", "2026-07-01", 16000, 18500, 12, ""),
    ("青眼の白龍 初期（遊戯王）", "遊戯王", "青眼の白龍", "", "初期弾",
     "○", "×", "SS", "安定", "低", "在庫中", "2026-05-20", 38000, 49000, 55, "絶版・入手困難"),
    ("ピカチュウ プロモ（ポケモンカード）", "ポケモンカード", "ピカチュウ", "サンプル絵師B", "25周年記念",
     "×", "○", "A", "安定", "中", "受注済", "2026-07-10", 7000, 8300, 40, "周年記念"),
    ("一般パック品（遊戯王）", "遊戯王", "", "", "",
     "×", "×", "C", "下降", "高", "売却済", "2026-04-01", 2800, 2900, 90, ""),
]
name_cols = ["商品名", "銘柄タイトル", "キャラ", "イラストレーター", "企画・収録弾", "絶版", "周年",
             "希少性", "トレンド", "ボラティリティ", "ステータス", "仕入日", "仕入値",
             "現在時価", "想定売却日数", "メモ"]
for i, values in enumerate(inv_samples):
    r = INV_DATA_START + i
    for name, val in zip(name_cols, values):
        ws[f"{IC[name]}{r}"] = val

for r in range(INV_DATA_START, INV_DATA_END + 1):
    B = IC["商品名"]
    Kc, Nc, Oc, Pc, Qc, Rc, Sc = (IC["仕入値"], IC["現在時価"], IC["想定売却日数"], IC["想定粗利額"],
                                  IC["想定粗利率"], IC["象限"], IC["ポートフォリオ比率"])
    Vc, Eff, Cap, Pos = IC["ボラティリティ"], IC["実効想定売却日数"], IC["資本効率"], IC["ポジションサイズ目安"]
    ws[f"{IC['No']}{r}"] = f'=IF({B}{r}="","",ROW()-{INV_DATA_START-1})'
    ws[f"{Pc}{r}"] = f'=IF({B}{r}="","",{Nc}{r}-{Kc}{r})'
    ws[f"{Qc}{r}"] = f'=IF({B}{r}="","",IF({Kc}{r}=0,"",{Pc}{r}/{Kc}{r}))'
    ws[f"{Eff}{r}"] = (f'=IF({B}{r}="","",{Oc}{r}*IFERROR(INDEX({VOL_FACTOR_RANGE},MATCH({Vc}{r},{VOL_RANGE},0)),1))')
    ws[f"{Cap}{r}"] = f'=IF({B}{r}="","",IF({Oc}{r}=0,"",{Qc}{r}*365/{Oc}{r}))'
    ws[f"{Pos}{r}"] = (f'=IF({B}{r}="","",IFERROR({Qc}{r}/INDEX({VOL_FACTOR_RANGE},MATCH({Vc}{r},{VOL_RANGE},0)),""))')
    ws[f"{Rc}{r}"] = (f'=IF({B}{r}="","",IF({Oc}{r}<=設定!$C$6,IF({Qc}{r}>=設定!$C$7,"I(コア)","II(回転)"),'
                       f'IF({Qc}{r}>=設定!$C$7,"III(厳選)","IV(見送り)")))')
    ws[f"{Sc}{r}"] = f'=IF({B}{r}="","",IF({ST}{r}="在庫中",IF($C$6=0,"",{Nc}{r}/$C$6),""))'
    col_font_fmt = {
        "No": (BLACK, INT), "商品名": (BLUE, None), "銘柄タイトル": (BLUE, None), "キャラ": (BLUE, None),
        "イラストレーター": (BLUE, None), "企画・収録弾": (BLUE, None), "絶版": (BLUE, None), "周年": (BLUE, None),
        "希少性": (BLUE, None), "トレンド": (BLUE, None), "ボラティリティ": (BLUE, None),
        "ステータス": (BLUE, None), "仕入日": (BLUE, DATE), "仕入値": (BLUE, CUR), "現在時価": (BLUE, CUR),
        "想定売却日数": (BLUE, '0"日"'), "実効想定売却日数": (BLACK, '0.0"日"'),
        "想定粗利額": (BLACK, CUR), "想定粗利率": (BLACK, PCT), "資本効率": (BLACK, PCT),
        "ポジションサイズ目安": (BLACK, "0.00"),
        "象限": (BLACK, None), "ポートフォリオ比率": (BLACK, PCT), "メモ": (BLUE, None),
    }
    for name, (font, fmt) in col_font_fmt.items():
        c = ws[f"{IC[name]}{r}"]
        c.font = font
        c.border = BORDER
        if fmt:
            c.number_format = fmt

dv_cat2 = DataValidation(type="list", formula1=CAT_RANGE, allow_blank=True)
dv_rarity = DataValidation(type="list", formula1=RARITY_RANGE, allow_blank=True)
dv_trend = DataValidation(type="list", formula1=TREND_RANGE, allow_blank=True)
dv_vol = DataValidation(type="list", formula1=VOL_RANGE, allow_blank=True)
dv_status = DataValidation(type="list", formula1=STATUS_RANGE, allow_blank=True)
dv_yn = DataValidation(type="list", formula1='"○,×"', allow_blank=True)
for dv in (dv_cat2, dv_rarity, dv_trend, dv_vol, dv_status, dv_yn):
    ws.add_data_validation(dv)
dv_cat2.add(f"{IC['銘柄タイトル']}{INV_DATA_START}:{IC['銘柄タイトル']}{INV_DATA_END}")
dv_rarity.add(f"{IC['希少性']}{INV_DATA_START}:{IC['希少性']}{INV_DATA_END}")
dv_trend.add(f"{IC['トレンド']}{INV_DATA_START}:{IC['トレンド']}{INV_DATA_END}")
dv_vol.add(f"{IC['ボラティリティ']}{INV_DATA_START}:{IC['ボラティリティ']}{INV_DATA_END}")
dv_status.add(f"{IC['ステータス']}{INV_DATA_START}:{IC['ステータス']}{INV_DATA_END}")
dv_yn.add(f"{IC['絶版']}{INV_DATA_START}:{IC['周年']}{INV_DATA_END}")

print("在庫管理 sheet done")

# =====================================================================
# 6. ダッシュボード
# =====================================================================
ws = sheet("ダッシュボード")
set_col_widths(ws, [3, 32, 16, 12, 12, 4, 20, 16, 12, 12, 12, 12])
title(ws, "B2", "ダッシュボード")
note(ws, "B3", "在庫管理シートを自動集計します。手入力は不要です（すべて数式）。")


def inv_range(colname):
    c = IC[colname]
    return f"在庫管理!${c}${INV_DATA_START}:${c}${INV_DATA_END}"


INV_L = inv_range("現在時価")
INV_I = inv_range("ステータス")
INV_M = inv_range("想定売却日数")
INV_P = inv_range("象限")
INV_C = inv_range("銘柄タイトル")
INV_D = inv_range("キャラ")
INV_E = inv_range("イラストレーター")


def kpi(ws, row, label, formula, fmt, link=False):
    ws.cell(row=row, column=2, value=label).font = BLACK
    c = ws.cell(row=row, column=3, value=formula)
    c.font = GREEN if link else BLACK
    c.number_format = fmt
    for col in (2, 3):
        ws.cell(row=row, column=col).border = BORDER


section(ws, "B5", "■ 資金効率")
kpi(ws, 6, "総仕入額（在庫中）", "=在庫管理!$C$5", CUR, link=True)
kpi(ws, 7, "総時価（在庫中）", "=在庫管理!$C$6", CUR, link=True)
kpi(ws, 8, "含み損益", "=在庫管理!$C$7", CUR, link=True)
kpi(ws, 9, "含み損益率", "=在庫管理!$C$8", PCT, link=True)
kpi(ws, 10, "資金倍率（時価 ÷ 仕入額）", '=IFERROR($C$7/$C$6,"")', "0.00x")
kpi(ws, 11, "平均想定売却日数（時価加重）", f'=IFERROR(SUMPRODUCT(({INV_I}="在庫中")*{INV_M}*{INV_L})/$C$7,"")', '0.0"日"')
kpi(ws, 12, "在庫の入れ替わり回数（年間・365日÷平均売却日数）", '=IFERROR(365/$C$11,"")', "0.00")

section(ws, "B14", "■ 換金のしやすさ")
kpi(ws, 15, "早く売れる基準以内の在庫比率（時価ベース）",
    f'=IFERROR(SUMIFS({INV_L},{INV_I},"在庫中",{INV_M},"<="&設定!$C$6)/$C$7,"")', PCT)
kpi(ws, 16, "売れ残りリスクが高い在庫比率（時価ベース）",
    f'=IFERROR(SUMIFS({INV_L},{INV_I},"在庫中",{INV_M},">"&設定!$C$8)/$C$7,"")', PCT)
kpi(ws, 17, "回転が速い在庫比率（コア＋回転、時価ベース）",
    f'=IFERROR((SUMIFS({INV_L},{INV_I},"在庫中",{INV_P},"I(コア)")+SUMIFS({INV_L},{INV_I},"在庫中",{INV_P},"II(回転)"))/$C$7,"")', PCT)
kpi(ws, 18, "回転が遅い在庫比率（厳選＋見送り、時価ベース）",
    f'=IFERROR((SUMIFS({INV_L},{INV_I},"在庫中",{INV_P},"III(厳選)")+SUMIFS({INV_L},{INV_I},"在庫中",{INV_P},"IV(見送り)"))/$C$7,"")', PCT)

section(ws, "B20", "■ 偏りリスク分析（特定の銘柄・カテゴリに偏っていないかを確認。数値が大きいほど偏りが大きい。目安：1,500未満=分散／1,500〜2,500=中程度／2,500以上=偏りが大きい）")


def conc_table_header(ws, row, col_start, name_label):
    header_row(ws, row, col_start, [name_label, "時価（在庫中）", "構成比", "偏りスコア"])


def conc_total_row(ws, row, col_start, data_start, data_end):
    L = get_column_letter(col_start + 1)
    S = get_column_letter(col_start + 2)
    H = get_column_letter(col_start + 3)
    ws.cell(row=row, column=col_start, value="合計").font = BOLD
    ws.cell(row=row, column=col_start + 1, value=f"=SUM({L}{data_start}:{L}{data_end})")
    ws.cell(row=row, column=col_start + 2, value=f"=SUM({S}{data_start}:{S}{data_end})")
    ws.cell(row=row, column=col_start + 3, value=f"=SUM({H}{data_start}:{H}{data_end})")
    for c in range(col_start, col_start + 4):
        cc = ws.cell(row=row, column=c)
        cc.font = BOLD
        cc.border = BORDER
    ws.cell(row=row, column=col_start + 1).number_format = CUR
    ws.cell(row=row, column=col_start + 2).number_format = PCT
    ws.cell(row=row, column=col_start + 3).number_format = "0"
    risk_row = row + 1
    Hcell = f"{H}{row}"
    ws.cell(row=risk_row, column=col_start, value="偏りの判定")
    ws.cell(row=risk_row, column=col_start).font = BOLD
    ws.cell(row=risk_row, column=col_start + 1,
            value=f'=IF({Hcell}>=2500,"偏りが大きい（要注意）",IF({Hcell}>=1500,"中程度","分散（健全）"))')
    ws.cell(row=risk_row, column=col_start + 1).font = BOLD
    ws.merge_cells(start_row=risk_row, start_column=col_start + 1, end_row=risk_row, end_column=col_start + 3)
    return risk_row


# --- 銘柄タイトル別（B列） ---
CAT_TBL_HDR = 22
CAT_TBL_START = 23
conc_table_header(ws, CAT_TBL_HDR, 2, "銘柄タイトル")
for i in range(6):
    r = CAT_TBL_START + i
    setrow = 20 + i
    ws.cell(row=r, column=2, value=f"='設定'!B{setrow}")
    ws.cell(row=r, column=3, value=f'=SUMIFS({INV_L},{INV_I},"在庫中",{INV_C},$B{r})')
    ws.cell(row=r, column=4, value=f'=IFERROR($C{r}/$C$7,"")')
    ws.cell(row=r, column=5, value=f'=IFERROR($D{r}^2*10000,0)')
    for col, fmt in [(2, None), (3, CUR), (4, PCT), (5, "0")]:
        c = ws.cell(row=r, column=col)
        c.font = BLACK
        c.border = BORDER
        if fmt:
            c.number_format = fmt
CAT_TBL_END = CAT_TBL_START + 5
cat_risk_row = conc_total_row(ws, CAT_TBL_END + 1, 2, CAT_TBL_START, CAT_TBL_END)

# --- 象限別（G列） ---
QUAD_TBL_HDR = 22
QUAD_TBL_START = 23
conc_table_header(ws, QUAD_TBL_HDR, 7, "象限")
quad_labels = ["I(コア)", "II(回転)", "III(厳選)", "IV(見送り)"]
for i, q in enumerate(quad_labels):
    r = QUAD_TBL_START + i
    ws.cell(row=r, column=7, value=q)
    ws.cell(row=r, column=8, value=f'=SUMIFS({INV_L},{INV_I},"在庫中",{INV_P},$G{r})')
    ws.cell(row=r, column=9, value=f'=IFERROR($H{r}/$C$7,"")')
    ws.cell(row=r, column=10, value=f'=IFERROR($I{r}^2*10000,0)')
    for col, fmt in [(7, None), (8, CUR), (9, PCT), (10, "0")]:
        c = ws.cell(row=r, column=col)
        c.font = BLACK
        c.border = BORDER
        if fmt:
            c.number_format = fmt
QUAD_TBL_END = QUAD_TBL_START + 3
quad_risk_row = conc_total_row(ws, QUAD_TBL_END + 1, 7, QUAD_TBL_START, QUAD_TBL_END)

# --- キャラ別（B列、その他行含む） ---
CHAR_TBL_HDR = 32
CHAR_TBL_START = 33
conc_table_header(ws, CHAR_TBL_HDR, 2, "キャラ")
for i in range(20):
    r = CHAR_TBL_START + i
    setrow = CHAR_START + i
    ws.cell(row=r, column=2, value=f"='設定'!B{setrow}")
    ws.cell(row=r, column=3, value=f'=IF($B{r}="","",SUMIFS({INV_L},{INV_I},"在庫中",{INV_D},$B{r}))')
    ws.cell(row=r, column=4, value=f'=IFERROR($C{r}/$C$7,"")')
    ws.cell(row=r, column=5, value=f'=IFERROR($D{r}^2*10000,0)')
    for col, fmt in [(2, None), (3, CUR), (4, PCT), (5, "0")]:
        c = ws.cell(row=r, column=col)
        c.font = BLACK
        c.border = BORDER
        if fmt:
            c.number_format = fmt
OTHER_CHAR_ROW = CHAR_TBL_START + 20
ws.cell(row=OTHER_CHAR_ROW, column=2, value="その他（未登録キャラ／タグなし）")
ws.cell(row=OTHER_CHAR_ROW, column=3,
        value=f'=MAX(0,SUMIFS({INV_L},{INV_I},"在庫中")-SUM(C{CHAR_TBL_START}:C{CHAR_TBL_START+19}))')
ws.cell(row=OTHER_CHAR_ROW, column=4, value=f'=IFERROR($C{OTHER_CHAR_ROW}/$C$7,"")')
ws.cell(row=OTHER_CHAR_ROW, column=5, value=f'=IFERROR($D{OTHER_CHAR_ROW}^2*10000,0)')
for col, fmt in [(2, None), (3, CUR), (4, PCT), (5, "0")]:
    c = ws.cell(row=OTHER_CHAR_ROW, column=col)
    c.font = NOTE
    c.border = BORDER
    if fmt:
        c.number_format = fmt
char_risk_row = conc_total_row(ws, OTHER_CHAR_ROW + 1, 2, CHAR_TBL_START, OTHER_CHAR_ROW)

# --- イラストレーター別（G列、その他行含む） ---
ILL_TBL_HDR = 32
ILL_TBL_START = 33
conc_table_header(ws, ILL_TBL_HDR, 7, "イラストレーター")
for i in range(20):
    r = ILL_TBL_START + i
    setrow = ILLUST_START + i
    ws.cell(row=r, column=7, value=f"='設定'!B{setrow}")
    ws.cell(row=r, column=8, value=f'=IF($G{r}="","",SUMIFS({INV_L},{INV_I},"在庫中",{INV_E},$G{r}))')
    ws.cell(row=r, column=9, value=f'=IFERROR($H{r}/$C$7,"")')
    ws.cell(row=r, column=10, value=f'=IFERROR($I{r}^2*10000,0)')
    for col, fmt in [(7, None), (8, CUR), (9, PCT), (10, "0")]:
        c = ws.cell(row=r, column=col)
        c.font = BLACK
        c.border = BORDER
        if fmt:
            c.number_format = fmt
OTHER_ILL_ROW = ILL_TBL_START + 20
ws.cell(row=OTHER_ILL_ROW, column=7, value="その他（未登録／タグなし）")
ws.cell(row=OTHER_ILL_ROW, column=8,
        value=f'=MAX(0,SUMIFS({INV_L},{INV_I},"在庫中")-SUM(H{ILL_TBL_START}:H{ILL_TBL_START+19}))')
ws.cell(row=OTHER_ILL_ROW, column=9, value=f'=IFERROR($H{OTHER_ILL_ROW}/$C$7,"")')
ws.cell(row=OTHER_ILL_ROW, column=10, value=f'=IFERROR($I{OTHER_ILL_ROW}^2*10000,0)')
for col, fmt in [(7, None), (8, CUR), (9, PCT), (10, "0")]:
    c = ws.cell(row=OTHER_ILL_ROW, column=col)
    c.font = NOTE
    c.border = BORDER
    if fmt:
        c.number_format = fmt
ill_risk_row = conc_total_row(ws, OTHER_ILL_ROW + 1, 7, ILL_TBL_START, OTHER_ILL_ROW)

# --- 攻めと守りの評価（利益率×回転率 vs 流動性×ボラティリティ） ---
section(ws, "B58", "■ 攻めと守りの評価（資本効率・実効流動性）")
note(ws, "B59", "利益率×回転率＝「どれだけ稼ぐ力があるか」（攻めの評価・仕入れ優先順位の決定に使用）。"
                "流動性×ボラティリティ＝「いざという時に安全に現金化できるか」（守りの評価・ポジションサイズの抑制判断に使用）。")
kpi(ws, 60, "平均資本効率（利益率×回転率・時価加重）",
    f'=IFERROR(SUMPRODUCT(({INV_I}="在庫中")*N({inv_range("資本効率")})*N({INV_L}))/$C$7,"")', PCT)
kpi(ws, 61, "平均実効想定売却日数（ボラティリティ考慮・時価加重）",
    f'=IFERROR(SUMPRODUCT(({INV_I}="在庫中")*N({inv_range("実効想定売却日数")})*N({INV_L}))/$C$7,"")', '0.0"日"')

# --- 象限別 枚数構成比（分布傾向のチェック） ---
section(ws, "B63", "■ 象限別 枚数構成比（分布傾向のチェック）")
note(ws, "B64", "時価ではなく点数（枚数）ベースの構成比。一般的には「コアが薄い層、回転（キャッシュ同等物）が最大ボリューム、"
                "厳選（グロース枠）が最少数・最高ボラ、見送り（劣後在庫）は低ボラで少量」という逆ピラミッド型になりやすい。")
header_row(ws, 65, 2, ["象限", "点数（在庫中）", "構成比"])
QCNT_START = 66
for i, q in enumerate(quad_labels):
    r = QCNT_START + i
    ws.cell(row=r, column=2, value=q)
    ws.cell(row=r, column=3, value=f'=COUNTIFS({INV_I},"在庫中",{INV_P},$B{r})')
    ws.cell(row=r, column=4, value=f'=IFERROR($C{r}/$C${QCNT_START+4},"")')
    for col, fmt in [(2, None), (3, INT), (4, PCT)]:
        c = ws.cell(row=r, column=col)
        c.font = BLACK
        c.border = BORDER
        if fmt:
            c.number_format = fmt
QCNT_END = QCNT_START + 3
ws.cell(row=QCNT_END + 1, column=2, value="合計").font = BOLD
ws.cell(row=QCNT_END + 1, column=3, value=f"=SUM(C{QCNT_START}:C{QCNT_END})")
ws.cell(row=QCNT_END + 1, column=3).number_format = INT
ws.cell(row=QCNT_END + 1, column=3).font = BOLD
ws.cell(row=QCNT_END + 1, column=3).border = BORDER
ws.cell(row=QCNT_END + 1, column=2).border = BORDER
ws.cell(row=QCNT_END + 1, column=4, value=f"=SUM(D{QCNT_START}:D{QCNT_END})")
ws.cell(row=QCNT_END + 1, column=4).number_format = PCT
ws.cell(row=QCNT_END + 1, column=4).font = BOLD
ws.cell(row=QCNT_END + 1, column=4).border = BORDER

# --- 希少性帯別 偏りリスク（レアリティ帯単位の相関ドライバー） ---
RARITY_TBL_HDR = 72
RARITY_TBL_START = 73
section(ws, f"B{RARITY_TBL_HDR-1}", "■ 希少性帯別・企画別 偏りリスク（相関ドライバー：収録弾・レアリティ帯・企画単位の集中を確認）")
conc_table_header(ws, RARITY_TBL_HDR, 2, "希少性")
for i in range(5):
    r = RARITY_TBL_START + i
    setrow = 29 + i
    ws.cell(row=r, column=2, value=f"='設定'!B{setrow}")
    ws.cell(row=r, column=3, value=f'=SUMIFS({INV_L},{INV_I},"在庫中",在庫管理!${IC["希少性"]}${INV_DATA_START}:${IC["希少性"]}${INV_DATA_END},$B{r})')
    ws.cell(row=r, column=4, value=f'=IFERROR($C{r}/$C$7,"")')
    ws.cell(row=r, column=5, value=f'=IFERROR($D{r}^2*10000,0)')
    for col, fmt in [(2, None), (3, CUR), (4, PCT), (5, "0")]:
        c = ws.cell(row=r, column=col)
        c.font = BLACK
        c.border = BORDER
        if fmt:
            c.number_format = fmt
RARITY_TBL_END = RARITY_TBL_START + 4
rarity_risk_row = conc_total_row(ws, RARITY_TBL_END + 1, 2, RARITY_TBL_START, RARITY_TBL_END)

# --- 企画・収録弾別 偏りリスク（相関ドライバー、その他行含む） ---
CAMP_TBL_HDR = RARITY_TBL_HDR
CAMP_TBL_START = RARITY_TBL_START
conc_table_header(ws, CAMP_TBL_HDR, 7, "企画・収録弾")
INV_CAMP = inv_range("企画・収録弾")
for i in range(20):
    r = CAMP_TBL_START + i
    setrow = CAMPAIGN_START + i
    ws.cell(row=r, column=7, value=f"='設定'!B{setrow}")
    ws.cell(row=r, column=8, value=f'=IF($G{r}="","",SUMIFS({INV_L},{INV_I},"在庫中",{INV_CAMP},$G{r}))')
    ws.cell(row=r, column=9, value=f'=IFERROR($H{r}/$C$7,"")')
    ws.cell(row=r, column=10, value=f'=IFERROR($I{r}^2*10000,0)')
    for col, fmt in [(7, None), (8, CUR), (9, PCT), (10, "0")]:
        c = ws.cell(row=r, column=col)
        c.font = BLACK
        c.border = BORDER
        if fmt:
            c.number_format = fmt
OTHER_CAMP_ROW = CAMP_TBL_START + 20
ws.cell(row=OTHER_CAMP_ROW, column=7, value="その他（未登録／タグなし）")
ws.cell(row=OTHER_CAMP_ROW, column=8,
        value=f'=MAX(0,SUMIFS({INV_L},{INV_I},"在庫中")-SUM(H{CAMP_TBL_START}:H{CAMP_TBL_START+19}))')
ws.cell(row=OTHER_CAMP_ROW, column=9, value=f'=IFERROR($H{OTHER_CAMP_ROW}/$C$7,"")')
ws.cell(row=OTHER_CAMP_ROW, column=10, value=f'=IFERROR($I{OTHER_CAMP_ROW}^2*10000,0)')
for col, fmt in [(7, None), (8, CUR), (9, PCT), (10, "0")]:
    c = ws.cell(row=OTHER_CAMP_ROW, column=col)
    c.font = NOTE
    c.border = BORDER
    if fmt:
        c.number_format = fmt
camp_risk_row = conc_total_row(ws, OTHER_CAMP_ROW + 1, 7, CAMP_TBL_START, OTHER_CAMP_ROW)

print("ダッシュボード 追加テーブル done, risk rows:", rarity_risk_row, camp_risk_row)

# 3-color scale conditional formatting on 偏りスコア columns
for col_letter, start, end in [("E", CAT_TBL_START, CAT_TBL_END), ("J", QUAD_TBL_START, QUAD_TBL_END),
                                ("E", CHAR_TBL_START, OTHER_CHAR_ROW), ("J", ILL_TBL_START, OTHER_ILL_ROW),
                                ("E", RARITY_TBL_START, RARITY_TBL_END), ("J", CAMP_TBL_START, OTHER_CAMP_ROW)]:
    rng = f"{col_letter}{start}:{col_letter}{end}"
    ws.conditional_formatting.add(
        rng,
        ColorScaleRule(start_type="num", start_value=0, start_color="C6EFCE",
                        mid_type="num", mid_value=2500, mid_color="FFEB9C",
                        end_type="num", end_value=10000, end_color="FFC7CE"),
    )

print("ダッシュボード 偏りリスク table done, risk rows:", cat_risk_row, quad_risk_row, char_risk_row, ill_risk_row)

# --- charts ---
pie = PieChart()
pie.title = "ポートフォリオ構成（銘柄タイトル別・時価）"
pie_data = Reference(ws, min_col=3, min_row=CAT_TBL_HDR, max_row=CAT_TBL_END)
pie_cats = Reference(ws, min_col=2, min_row=CAT_TBL_START, max_row=CAT_TBL_END)
pie.add_data(pie_data, titles_from_data=True)
pie.set_categories(pie_cats)
pie.height, pie.width = 8, 14
ws.add_chart(pie, "M5")

bar = BarChart()
bar.type = "col"
bar.title = "象限別 時価（在庫中）"
bar.y_axis.title = "時価（円）"
bar_data = Reference(ws, min_col=8, min_row=QUAD_TBL_HDR, max_row=QUAD_TBL_END)
bar_cats = Reference(ws, min_col=7, min_row=QUAD_TBL_START, max_row=QUAD_TBL_END)
bar.add_data(bar_data, titles_from_data=True)
bar.set_categories(bar_cats)
bar.height, bar.width = 8, 14
bar.legend = None
ws.add_chart(bar, "M22")

print("ダッシュボード charts done")

# =====================================================================
# final touches
# =====================================================================
tab_colors = {
    "使い方": "808080", "設定": "808080", "ウォッチリスト": "2E75B6", "仕入れ計画": "548235",
    "在庫管理": "BF8F00", "ダッシュボード": "C00000",
}
for name, color in tab_colors.items():
    wb[name].sheet_properties.tabColor = color

for name in wb.sheetnames:
    s = wb[name]
    s.page_setup.orientation = "landscape"
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 0
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.page_margins.left = s.page_margins.right = 0.4
    s.page_margins.top = s.page_margins.bottom = 0.5

wb["ウォッチリスト"].print_title_rows = f"{WL_HEADER_ROW}:{WL_HEADER_ROW}"
wb["在庫管理"].print_title_rows = f"{INV_HEADER_ROW}:{INV_HEADER_ROW}"

wb.active = wb.sheetnames.index("使い方")
wb.save(OUT_PATH)
print("ALL SHEETS SAVED")
